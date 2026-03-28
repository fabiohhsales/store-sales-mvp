from pathlib import Path
from docx import Document
from docx.shared import Pt


def add_heading(doc: Document, text: str, level: int = 1):
    doc.add_heading(text, level=level)


def add_paragraph(doc: Document, text: str, bold: bool = False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    return p


def add_bullet(doc: Document, text: str):
    doc.add_paragraph(text, style='List Bullet')


def add_number(doc: Document, text: str):
    doc.add_paragraph(text, style='List Number')


def main() -> None:
    out_path = Path('c:/Users/Pichau/panel_novo/painel2/docs/FollowUps_Logica_Atual_ChatSales.docx')
    out_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()

    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    add_heading(doc, 'Follow Ups - Logica Atual no Painel2 ChatSales', level=0)
    add_paragraph(doc, 'Documento tecnico-funcional com explicacao simples da logica atual de Follow Ups, desde a entrada da mensagem ate os controles de envio.')
    add_paragraph(doc, 'Data: 27/03/2026')

    add_heading(doc, '1. Resumo Executivo (bem simples)', level=1)
    add_number(doc, 'O paciente manda mensagem no Chatwoot.')
    add_number(doc, 'O webhook recebe, valida e salva a mensagem.')
    add_number(doc, 'O sistema marca quando o paciente falou por ultimo (last_incoming_at).')
    add_number(doc, 'Se a IA responder, marca quando o bot falou por ultimo (last_outgoing_at).')
    add_number(doc, 'Crons rodam periodicamente e checam se chegou a hora de algum follow-up.')
    add_number(doc, 'Antes de enviar, o sistema verifica controle de duplicidade (idempotencia).')
    add_number(doc, 'Se ja enviou aquele passo para aquela conversa, nao envia de novo.')
    add_number(doc, 'Quando envia, registra log e atualiza campos de acompanhamento.')

    add_heading(doc, '2. Entrada e Pipeline Inicial', level=1)
    add_bullet(doc, 'Webhook principal: src/app/api/webhooks/chatwoot/route.ts')
    add_bullet(doc, 'Aceita evento message_created e considera incoming (paciente).')
    add_bullet(doc, 'Ignora mensagens privadas e grupos (identifier terminando em @g.us).')
    add_bullet(doc, 'Normaliza payload e dispara pipeline em background.')
    add_bullet(doc, 'Pipeline base: src/lib/bot/pipeline.ts')
    add_bullet(doc, 'Cria/atualiza contato e conversa, salva mensagem e historico.')
    add_bullet(doc, 'Atualiza last_incoming_at quando chega mensagem do paciente.')

    add_heading(doc, '3. Quando a IA responde', level=1)
    add_bullet(doc, 'Dispatcher: src/lib/bot/dispatcher.ts')
    add_bullet(doc, 'Envia resposta via Evolution API, salva mensagem da IA e atualiza Chatwoot.')
    add_bullet(doc, 'Atualiza labels da conversa e last_outgoing_at (quando houve reply).')
    add_bullet(doc, 'Com isso, o sistema sabe quem falou por ultimo.')

    add_heading(doc, '4. Motores de Follow Up existentes hoje', level=1)
    add_paragraph(doc, 'Hoje existem duas trilhas para agendado (legado e nova) e duas cadencias novas para funil (lead e atendimento).', bold=True)

    add_heading(doc, '4.1 Legado de Agendado (confirmacao, lembrete, no-show)', level=2)
    add_bullet(doc, 'Arquivo: src/lib/followup/confirmations.ts')
    add_bullet(doc, 'Usa followup_confirmation_hours_before (padrao 24h) para confirmacao.')
    add_bullet(doc, 'Usa followup_reminder_hours_before (padrao 2h) para lembrete.')
    add_bullet(doc, 'No-show: consulta passou, tinha confirmacao enviada e ainda esta na janela recente.')
    add_bullet(doc, 'Controle de repeticao principal em campos da tabela appointments: confirmation_sent_at e reminder_sent_at.')

    add_heading(doc, '4.2 Cadencia Lead', level=2)
    add_bullet(doc, 'Arquivo: src/lib/followup/lead-cadence.ts')
    add_bullet(doc, 'Entra quando bot falou por ultimo: last_incoming_at < last_outgoing_at.')
    add_bullet(doc, 'Nao entra se a conversa ja tem appointment futuro com status scheduled.')
    add_bullet(doc, 'Steps por janela de horas desde o ultimo outgoing: D+1, D+2, D+3, D+5, D+7.')
    add_bullet(doc, 'Templates customizaveis por step no panel_bot_config.')

    add_heading(doc, '4.3 Cadencia Em Atendimento', level=2)
    add_bullet(doc, 'Arquivo: src/lib/followup/atendimento-cadence.ts')
    add_bullet(doc, 'Entra quando paciente falou por ultimo: last_incoming_at > last_outgoing_at (ou sem outgoing).')
    add_bullet(doc, 'Tambem ignora conversas com appointment futuro scheduled.')
    add_bullet(doc, 'Steps por janela de horas: D+1, D+2, D+4, D+7, D+10.')
    add_bullet(doc, 'Templates customizaveis por step no panel_bot_config.')

    add_heading(doc, '4.4 Cadencia Agendado (nova)', level=2)
    add_bullet(doc, 'Arquivo: src/lib/followup/agendado-cadence.ts')
    add_bullet(doc, 'Steps: D-2 12h, -3h, -5min antes do start_at.')
    add_bullet(doc, 'Calcula horas faltantes para consulta e decide step da janela ativa.')
    add_bullet(doc, 'Usa templates novos agendado_followup_msg_d2, agendado_followup_msg_minus3h e agendado_followup_msg_minus5min.')

    add_heading(doc, '5. Idempotencia e controles anti-duplicidade', level=1)
    add_bullet(doc, 'Migration: supabase/migrations/004_followup_cadence.sql')
    add_bullet(doc, 'Tabela followup_cadence_steps com UNIQUE(conversation_id, cadence_type, step_key).')
    add_bullet(doc, 'Nas cadencias novas, primeiro reserva o step com insert ignoreDuplicates.')
    add_bullet(doc, 'Se ja existir, significa que aquele passo ja foi enviado e o envio e pulado.')
    add_bullet(doc, 'Se falhar apos reservar (ex.: erro no envio), a reserva e removida para retry no proximo cron.')

    add_heading(doc, '6. Regras de horario', level=1)
    add_bullet(doc, 'Funcao base de horario comercial: 08:00-17:00 (America/Sao_Paulo).')
    add_bullet(doc, 'Lead e Atendimento param fora do horario comercial.')
    add_bullet(doc, 'Legado de confirmacoes tambem para fora do horario comercial.')
    add_bullet(doc, 'Agendado novo registra skippedOutsideHours, mas ainda processa os steps.')

    add_heading(doc, '7. Endpoints de cron que disparam os motores', level=1)
    add_bullet(doc, 'Legado: src/app/api/cron/confirmacoes/route.ts -> runFollowupPipeline().')
    add_bullet(doc, 'Cadencias Lead+Atendimento: src/app/api/cron/followup-cadencia/route.ts -> runLeadCadencePipeline() + runAtendimentoCadencePipeline().')
    add_bullet(doc, 'Agendado novo: src/app/api/cron/followup-agendado/route.ts -> runAgendadoCadencePipeline().')
    add_bullet(doc, 'Todos validam Authorization Bearer com CRON_SECRET.')

    add_heading(doc, '8. Configuracao no painel (UI e API)', level=1)
    add_bullet(doc, 'UI: src/components/bot-config/followup-section.tsx (abas Agendado, Lead, Em Atendimento).')
    add_bullet(doc, 'API de salvamento: src/app/api/bot-config/route.ts (inclui campos novos).')
    add_bullet(doc, 'Validacao: src/lib/validations/bot-config.ts (Zod).')
    add_bullet(doc, 'Tipos: src/types/database.ts.')

    add_heading(doc, '9. Logs e rastreabilidade', level=1)
    add_bullet(doc, 'Tabela followup_logs recebe registro do que foi enviado, quando e para qual conversa/contato.')
    add_bullet(doc, 'Campos de conversa atualizados nas cadencias novas: followup_cadence e last_followup_at.')

    add_heading(doc, '10. Ponto de atencao operacional', level=1)
    add_paragraph(doc, 'Atualmente coexistem dois motores para agendado (legado e novo). Se ambos os crons estiverem ativos ao mesmo tempo, ha sobreposicao operacional. Recomenda-se definir estrategia de transicao para evitar comportamentos concorrentes.')

    add_heading(doc, '11. Checklist rapido de verificacao em producao', level=1)
    add_number(doc, 'Confirmar CRON_SECRET configurado.')
    add_number(doc, 'Confirmar agendas dos crons (hourly e 5min) conforme estrategia escolhida.')
    add_number(doc, 'Verificar toggles e templates no panel_bot_config.')
    add_number(doc, 'Monitorar followup_logs e followup_cadence_steps apos ativacao.')
    add_number(doc, 'Validar se last_incoming_at e last_outgoing_at estao sendo atualizados corretamente.')

    doc.save(out_path)
    print(str(out_path))


if __name__ == '__main__':
    main()
